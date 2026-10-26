import io
import csv
from typing import List, Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from app.models.models import Lead, WebsiteAudit, OutreachDraft

INJECTION_CHARS = ("=", "+", "-", "@", "\t", "\r")


def sanitize_formula_injection(value: Any) -> Any:
    """
    Guards against CSV and spreadsheet formula injection.
    Prefixes values starting with '=', '+', '-', '@', '\t', '\r' with a single quote.
    """
    if isinstance(value, str):
        if value.startswith(INJECTION_CHARS):
            return "'" + value
    return value


def sanitize_dict_row(row: Dict[str, Any]) -> Dict[str, Any]:
    """Sanitizes all string fields in a dictionary row."""
    return {k: sanitize_formula_injection(v) for k, v in row.items()}


def leads_to_export_rows(leads: List[Lead]) -> List[Dict[str, Any]]:
    """Converts a list of Lead models to export-ready dictionaries with verification status and sources."""
    rows = []
    for lead in leads:
        row = {
            "Lead ID": lead.id,
            "Company Name": lead.company_name,
            "Official Website": lead.official_website or "",
            "Industry": lead.industry or "",
            "City": lead.city or "",
            "State": lead.state or "",
            "Country": lead.country or "",
            "Public Phone": lead.public_business_phone or "",
            "Phone Verification Status": lead.phone_verification_status or "Missing",
            "Business Email": lead.general_business_email or "",
            "Email Verification Status": lead.email_verification_status or "Missing",
            "Contact Name": lead.publicly_listed_contact_name or "",
            "Designation": lead.publicly_listed_designation or "",
            "LinkedIn URL": lead.company_linkedin_url or "",
            "Lead Score": lead.lead_score,
            "Priority": lead.priority,
            "CRM Status": lead.crm_status,
            "Marketing Opportunity": lead.marketing_opportunity or "",
            "Recommended Service": lead.recommended_service or "",
            "Data Source": lead.data_source or "",
            "Source URL": lead.source_url or "",
            "Last Verified Date": lead.last_verified_date.strftime("%Y-%m-%d %H:%M") if lead.last_verified_date else "",
            "Notes": lead.notes or "",
            "Created Date": lead.created_at.strftime("%Y-%m-%d %H:%M") if lead.created_at else ""
        }
        rows.append(sanitize_dict_row(row))
    return rows


def export_leads_csv(leads: List[Lead]) -> io.BytesIO:
    """Generates an injection-safe CSV file stream."""
    rows = leads_to_export_rows(leads)
    output = io.StringIO()
    if rows:
        writer = csv.DictWriter(output, fieldnames=list(rows[0].keys()), quoting=csv.QUOTE_MINIMAL)
        writer.writeheader()
        writer.writerows(rows)
    else:
        writer = csv.writer(output)
        writer.writerow(["No leads found"])
    
    mem = io.BytesIO()
    mem.write(output.getvalue().encode("utf-8-sig"))  # UTF-8 with BOM for Excel compatibility
    mem.seek(0)
    return mem


def export_leads_excel(leads: List[Lead], sheet_title: str = "Leads") -> io.BytesIO:
    """Generates a professional, injection-safe Excel (.xlsx) file stream."""
    rows = leads_to_export_rows(leads)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title

    # Styling
    header_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")  # Brand Blue
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    regular_font = Font(name="Calibri", size=10)
    border_side = Side(style="thin", color="E2E8F0")
    thin_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

    if rows:
        headers = list(rows[0].keys())
        ws.append(headers)

        # Style header row
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)

        # Add data rows
        for row_data in rows:
            row_values = list(row_data.values())
            ws.append(row_values)

        # Style data cells and auto-adjust column width
        for row in ws.iter_rows(min_row=2, max_row=len(rows) + 1, min_col=1, max_col=len(headers)):
            for cell in row:
                cell.font = regular_font
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)
    else:
        ws.append(["No records available for export"])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
